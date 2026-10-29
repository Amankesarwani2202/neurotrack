from flask import Flask, request, jsonify, render_template, send_file
import numpy as np
import io
import json
import os
import tempfile
import base64
import urllib.error
import urllib.parse
import urllib.request
from datetime import datetime, timezone
from openpyxl import Workbook
from openpyxl.utils import get_column_letter
from reportlab.lib.pagesizes import letter
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle

app = Flask(__name__)
DEFAULT_DATA_FILE = '/tmp/participants.json' if os.environ.get('VERCEL') else os.path.join(os.path.dirname(__file__), 'participants.json')
DATA_FILE = os.environ.get('NEUROTRACK_DATA_FILE', DEFAULT_DATA_FILE)
GITHUB_TOKEN = os.environ.get('GITHUB_TOKEN')
GITHUB_REPOSITORY = os.environ.get('GITHUB_REPOSITORY')
GITHUB_DATA_PATH = os.environ.get('GITHUB_DATA_PATH', 'participants.json')


def github_storage_enabled():
    return bool(GITHUB_TOKEN and GITHUB_REPOSITORY)


def github_file_url():
    path = urllib.parse.quote(GITHUB_DATA_PATH, safe='/')
    return f'https://api.github.com/repos/{GITHUB_REPOSITORY}/contents/{path}'


def github_request(method='GET', payload=None):
    body = json.dumps(payload).encode('utf-8') if payload is not None else None
    request = urllib.request.Request(
        github_file_url(),
        data=body,
        headers={
            'Accept': 'application/vnd.github+json',
            'Authorization': f'Bearer {GITHUB_TOKEN}',
            'X-GitHub-Api-Version': '2022-11-28'
        },
        method=method
    )
    with urllib.request.urlopen(request, timeout=15) as response:
        return json.loads(response.read().decode('utf-8'))


def load_github_participants():
    try:
        remote_file = github_request()
    except urllib.error.HTTPError as error:
        if error.code == 404:
            return [], None
        raise
    content = base64.b64decode(remote_file['content']).decode('utf-8')
    return json.loads(content), remote_file['sha']


def save_github_participants(participants, sha=None):
    # Retry optimistic-concurrency conflicts rather than losing concurrent writes.
    candidate, current_sha = participants, sha
    for attempt in range(4):
        if attempt:
            remote, current_sha = load_github_participants()
            merged = {item.get('id'): item for item in remote if item.get('id') is not None}
            for item in candidate:
                item_id = item.get('id')
                if item_id is None:
                    continue
                existing = merged.get(item_id)
                if existing is None:
                    merged[item_id] = item
                    continue
                existing_key = existing.get('idempotency_key')
                candidate_key = item.get('idempotency_key')
                if existing_key and candidate_key and existing_key != candidate_key:
                    # Two simultaneous registrations allocated the same next ID.
                    # Refuse to overwrite either record; caller can retry against fresh data.
                    raise RuntimeError('Concurrent participant registration detected; please retry')
                candidate_updated = item.get('analyzed_at') or ''
                existing_updated = existing.get('analyzed_at') or ''
                if candidate_updated and candidate_updated > existing_updated:
                    merged[item_id] = item
            candidate = list(merged.values())
        payload = {
            'message': 'Update participant data',
            'content': base64.b64encode((json.dumps(candidate, indent=2) + '\n').encode('utf-8')).decode('ascii')
        }
        if current_sha:
            payload['sha'] = current_sha
        try:
            github_request(method='PUT', payload=payload)
            return
        except urllib.error.HTTPError as error:
            if error.code not in (409, 422) or attempt == 3:
                raise
    raise RuntimeError('Could not save participant data after concurrency retries')

def load_participants():
    if github_storage_enabled():
        participants, _ = load_github_participants()
        return participants
    try:
        with open(DATA_FILE, 'r', encoding='utf-8') as data_file:
            return json.load(data_file)
    except (FileNotFoundError, json.JSONDecodeError):
        return []


def save_participants(participants):
    if github_storage_enabled():
        _, sha = load_github_participants()
        save_github_participants(participants, sha)
        return
    directory = os.path.dirname(DATA_FILE) or '.'
    os.makedirs(directory, exist_ok=True)
    with tempfile.NamedTemporaryFile('w', encoding='utf-8', dir=directory, delete=False) as temp_file:
        json.dump(participants, temp_file, indent=2)
        temp_file.write('\n')
        temporary_path = temp_file.name
    os.replace(temporary_path, DATA_FILE)

# --- Test Pages ---
@app.route('/')
def index():
    return render_template('index.html')

@app.route('/reaction')
def reaction():
    return render_template('reaction.html')

@app.route('/stroop')
def stroop():
    return render_template('stroop.html')

@app.route('/memory')
def memory():
    return render_template('memory.html')

@app.route('/clock')
def clock():
    return render_template('clock.html')

@app.route('/cookie-theft')
def cookie_theft():
    return render_template('cookie_theft.html')

@app.route('/letter-search')
def letter_search():
    return render_template('letter_search.html')

@app.route('/spiral')
def spiral():
    return render_template('spiral.html')

@app.route('/results')
def results():
    return render_template('results.html')

@app.route('/health')
def health():
    return jsonify({'status':'ok'})


@app.route('/participants', methods=['POST'])
def create_participant():
    data = request.json or {}
    required = ('name', 'medical_history', 'age', 'gender')
    if any(not data.get(field) for field in required):
        return jsonify({'error': 'name, medical_history, age, and gender are required'}), 400
    try:
        age = int(data['age'])
    except (TypeError, ValueError):
        return jsonify({'error': 'age must be a number'}), 400
    if age < 1 or age > 120:
        return jsonify({'error': 'age must be between 1 and 120'}), 400

    try:
        participants = load_participants()
    except (OSError, urllib.error.URLError, urllib.error.HTTPError, json.JSONDecodeError, RuntimeError) as error:
        return jsonify({'error': f'Participant storage is unavailable: {error}'}), 503
    idempotency_key = str(data.get('idempotency_key', '')).strip()
    if idempotency_key:
        existing = next((item for item in participants if item.get('idempotency_key') == idempotency_key), None)
        if existing:
            return jsonify({'participant_id': existing['id'], 'existing': True}), 200
    name = str(data['name']).strip()
    if not name:
        return jsonify({'error': 'name must not be blank'}), 400
    participant = {
        'id': max((item.get('id', 0) for item in participants if isinstance(item.get('id'), int)), default=0) + 1,
        'idempotency_key': idempotency_key or None,
        'name': name,
        'medical_history': data['medical_history'],
        'age': age,
        'gender': str(data['gender']).strip(),
        'test_scores': None,
        'created_at': datetime.now(timezone.utc).isoformat()
    }
    participants.append(participant)
    try:
        save_participants(participants)
    except (OSError, urllib.error.URLError, urllib.error.HTTPError, json.JSONDecodeError) as error:
        return jsonify({'error': f'Participant storage is unavailable: {error}'}), 503
    return jsonify({'participant_id': participant['id']}), 201


@app.route('/participants', methods=['GET'])
def list_participants():
    return jsonify(load_participants())

# Spiral analysis: expects points list [{x,y,t},...]
def analyze_spiral(points):
    if not points or len(points)<3:
        return {'spiral_score':0, 'notes':'too few points'}
    pts = np.array([[p['x'], p['y'], p.get('t', i*0.01)] for i,p in enumerate(points)])
    xy = pts[:,0:2]
    t = pts[:,2]
    dt = np.diff(t)
    dt[dt==0]=1e-6
    vel = np.linalg.norm(np.diff(xy,axis=0),axis=1)/dt
    if len(vel)==0:
        return {'spiral_score':0}
    jerk = np.mean(np.abs(np.diff(vel))) if len(vel)>1 else 0.0
    vel_std = float(np.std(vel))
    smoothness = 1.0/(1.0+vel_std)
    # deviation from simple polar linear fit
    xc,yc = np.mean(xy,axis=0)
    rel = xy - np.array([xc,yc])
    r = np.linalg.norm(rel,axis=1)
    theta = np.arctan2(rel[:,1], rel[:,0])
    try:
        A = np.vstack([theta, np.ones_like(theta)]).T
        a,b = np.linalg.lstsq(A, r, rcond=None)[0]
        r_fit = a*theta + b
        deviation = float(np.mean(np.abs(r - r_fit)))
    except Exception:
        deviation = float(np.mean(np.abs(r - np.mean(r))))
    dev_norm = np.tanh(deviation/50.0)
    jerk_norm = np.tanh(jerk/100.0)
    raw = (1 - dev_norm)*0.45 + (1 - jerk_norm)*0.35 + smoothness*0.2
    spiral_score = int(np.clip(raw,0,1)*100)
    return {'spiral_score': spiral_score, 'velocity_mean': float(np.mean(vel)), 'jerk': jerk, 'smoothness':smoothness, 'deviation': deviation}

# Main analyze_all endpoint
@app.route('/analyze_all', methods=['POST'])
def analyze_all():
    data = request.json or {}

    # Reaction: list of ms
    rts = data.get('reaction_times', [])
    rts = [float(value) for value in rts if isinstance(value, (int, float)) and np.isfinite(value) and value > 0]
    mean_rt = float(np.mean(rts)) if rts else None
    # map 150ms -> best(100), 1000ms -> worst(0)
    val = (mean_rt - 150)/(1000-150) if mean_rt is not None else 1
    val = np.clip(val,0,1)
    reaction_score = int((1-val)*100) if mean_rt is not None else None

    # Stroop: Use Avg RT from correct responses (lower is better)
    stroop_correct = max(0, int(data.get('stroop_correct', 0) or 0))
    stroop_total_raw = max(0, int(data.get('stroop_total', 0) or 0))
    stroop_total = stroop_total_raw or 1
    stroop_avg_rt_raw = data.get('stroop_avg_rt')
    stroop_avg_rt = float(stroop_avg_rt_raw) if isinstance(stroop_avg_rt_raw, (int, float)) and np.isfinite(stroop_avg_rt_raw) and stroop_avg_rt_raw > 0 else None

    # Combine accuracy and RT for Stroop score.
    acc_score = (stroop_correct/stroop_total) * 100
    rt_val = (stroop_avg_rt - 100)/(1000-100) if stroop_avg_rt is not None else 1
    rt_val = np.clip(rt_val, 0, 1)
    rt_score = (1-rt_val) * 100
    stroop_score = int(np.clip(acc_score * 0.3 + rt_score * 0.7, 0, 100)) if stroop_total_raw > 0 else None
    if stroop_total_raw > 0 and stroop_correct == 0:
        stroop_score = 0

    # Memory
    memory_correct = max(0, int(data.get('memory_correct', 0) or 0))
    memory_total_raw = max(0, int(data.get('memory_total', 0) or 0))
    memory_total = memory_total_raw or 1
    memory_score = int(np.clip(memory_correct / memory_total * 100, 0, 100)) if memory_total_raw > 0 else None

    # Clock drawing
    clock_data = data.get('clock', {}) if isinstance(data.get('clock', {}), dict) else {}
    clock_points_raw = clock_data.get('score_points')
    clock_points = max(0, min(6, int(clock_points_raw))) if clock_points_raw is not None else None
    clock_score = int(clock_points / 6.0 * 100) if clock_points is not None else None

    # Cookie theft
    cookie_data = data.get('cookie', {}) if isinstance(data.get('cookie', {}), dict) else {}
    cookie_score = int(np.clip(cookie_data['score'], 0, 100)) if isinstance(cookie_data.get('score'), (int, float)) else None

    # Letter search
    letter_data = data.get('letter', {}) if isinstance(data.get('letter', {}), dict) else {}
    letter_score = int(np.clip(letter_data['score'], 0, 100)) if isinstance(letter_data.get('score'), (int, float)) else (int(np.clip(letter_data['average_score'], 0, 100)) if isinstance(letter_data.get('average_score'), (int, float)) else None)

    # Spiral
    spiral_data = data.get('spiral', {})
    spiral_points = spiral_data.get('points', []) if isinstance(spiral_data, dict) else spiral_data
    if spiral_points and isinstance(spiral_points[0], dict) and 'points' in spiral_points[0]:
        spiral_points = [point for round_data in spiral_points for point in round_data.get('points', [])]
    spiral_res = analyze_spiral(spiral_points)
    spiral_score = spiral_res.get('spiral_score') if len(spiral_points) >= 5 else None
    history_proximity = {
        'self': 3,
        'close_relative': 2,
        'distant_relative': 1,
        'none': 0,
        'unknown': None
    }.get(data.get('medical_history'))
    expected_score_factor = {
        'self': 1.0,
        'close_relative': 0.75,
        'distant_relative': 0.5,
        'none': 0.0,
        'unknown': None
    }.get(data.get('medical_history'))

    # Combine with weights
    weights = {
        'reaction':1 / 7,
        'stroop':1 / 7,
        'memory':1 / 7,
        'spiral':1 / 7,
        'clock':1 / 7,
        'cookie':1 / 7,
        'letter':1 / 7
    }
    subs = {
        'reaction':reaction_score,
        'stroop':stroop_score,
        'memory':memory_score,
        'spiral':spiral_score,
        'clock':clock_score,
        'cookie':cookie_score,
        'letter':letter_score
    }
    available_scores = [score for score in subs.values() if isinstance(score, (int, float))]
    combined = float(np.mean(available_scores)) if available_scores else None
    risk_index = round(max(0, min(100, 100 - combined)), 2) if combined is not None else None

    out = {
        'user_name': data.get('user_name', 'Participant'),
        'medical_history': data.get('medical_history'),
        'history_proximity_score': history_proximity,
        'expected_score_factor': expected_score_factor,
        'age': data.get('age'),
        'gender': data.get('gender'),
        'subscores': {
            'Reaction_Time_Score': reaction_score,
            'Stroop_Score': stroop_score,
            'Memory_Score': memory_score,
            'Spiral_Score': spiral_score,
            'Clock_Drawing_Score': clock_score,
            'Cookie_Theft_Score': cookie_score,
            'Letter_Search_Score': letter_score
        },
        'raw_metrics': {
            'Reaction_Avg_ms': round(mean_rt, 2) if mean_rt is not None else None,
            'Stroop_Correct_Ratio': f"{stroop_correct}/{stroop_total_raw}",
            'Stroop_Avg_RT_ms': round(stroop_avg_rt, 2) if stroop_avg_rt is not None else None,
            'Memory_Correct_Ratio': f"{memory_correct}/{memory_total_raw}",
            'Clock_Drawing_Points': f"{clock_points}/6",
            'Cookie_Theft_Score': cookie_score,
            'Letter_Search_Score': letter_score
        },
        'spiral_details':spiral_res,
        'risk_index': risk_index
    }
    participant_id = data.get('participant_id')
    if participant_id:
        try:
            participants = load_participants()
            for participant in participants:
                if participant.get('id') == participant_id:
                    participant['test_scores'] = out['subscores']
                    participant['test_status'] = {key: ('completed' if value is not None else 'not_attempted') for key, value in out['subscores'].items()}
                    participant['raw_metrics'] = out['raw_metrics']
                    participant['spiral_details'] = out['spiral_details']
                    participant['risk_index'] = out['risk_index']
                    participant['history_proximity_score'] = out['history_proximity_score']
                    participant['expected_score_factor'] = out['expected_score_factor']
                    participant['analyzed_at'] = datetime.now(timezone.utc).isoformat()
                    break
            save_participants(participants)
        except (OSError, urllib.error.URLError, urllib.error.HTTPError, json.JSONDecodeError) as error:
            return jsonify({'error': f'Participant score storage is unavailable: {error}'}), 503
    return jsonify(out)

@app.route('/export.xlsx', methods=['GET'])
def export_participants_xlsx():
    try:
        participants = load_participants()
    except (OSError, urllib.error.URLError, urllib.error.HTTPError, json.JSONDecodeError) as error:
        return jsonify({'error': f'Participant storage is unavailable: {error}'}), 503
    workbook = Workbook()
    summary = workbook.active
    summary.title = 'Summary'
    summary.append(['Metric', 'Value'])
    summary.append(['Participants', len(participants)])
    summary.append(['Analyzed participants', sum(bool(item.get('test_scores')) for item in participants)])
    summary.append(['Exported at (UTC)', datetime.now(timezone.utc).isoformat()])
    sheet = workbook.create_sheet('Participants')
    columns = ['id', 'name', 'age', 'gender', 'medical_history', 'created_at', 'analyzed_at', 'test_status', 'test_scores', 'raw_metrics', 'spiral_details', 'risk_index']
    sheet.append(columns)
    for item in participants:
        sheet.append([json.dumps(item.get(key), ensure_ascii=False) if isinstance(item.get(key), (dict, list)) else item.get(key) for key in columns])
    audit = workbook.create_sheet('Data Quality')
    audit.append(['Issue', 'Details'])
    audit.append(['Missing vs zero', 'Blank score cells mean missing/not attempted; numeric zero means a recorded score.'])
    audit.append(['Legacy records', 'Legacy rows may not have test_status; do not infer attempted status from zero alone.'])
    audit.append(['Duplicate profiles', 'Potential duplicates are not automatically merged to avoid joining different people.'])
    for ws in workbook.worksheets:
        ws.freeze_panes = 'A2'
        ws.auto_filter.ref = ws.dimensions
        for col in ws.columns:
            width = min(max(max(len(str(cell.value or '')) for cell in col) + 2, 12), 48)
            ws.column_dimensions[get_column_letter(col[0].column)].width = width
    output = io.BytesIO()
    workbook.save(output)
    output.seek(0)
    return send_file(output, as_attachment=True, download_name='neurotrack_participants.xlsx',
                     mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')


# Generate PDF report
@app.route('/report', methods=['POST'])
def report():
    payload = request.json or {}
    results = payload.get('results', {})
    name = payload.get('name', 'Participant')

    descriptions = {
        'Reaction_Time_Score': 'Average time to respond. Higher is better (fast response).',
        'Stroop_Score': 'Cognitive flexibility and speed. Higher is better (accurate & fast).',
        'Memory_Score': 'Short-term numerical recall. Higher is better (correct recall).',
        'Clock_Drawing_Score': 'Planning and organisation of the clock drawing. Higher is better.',
        'Spiral_Score': 'Motor smoothness and tremor. Higher is better (smoother drawing).',
        'Cookie_Theft_Score': 'Visual scene interpretation. Higher is better.',
        'Letter_Search_Score': 'Visual search speed and accuracy. Higher is better.',
        'risk_index': 'Overall screening score (0=Lowest Risk, 100=Highest Risk). Lower is better.'
    }

    tmp = io.BytesIO()
    doc = SimpleDocTemplate(tmp, pagesize=letter, title=f"NEUROTRACK Report for {name}")
    styles = getSampleStyleSheet()
    styleSmall = ParagraphStyle('Small', parent=styles['Normal'], fontSize=8)

    story = [
        Paragraph("<b>NEUROTRACK - Cognitive & Motor Screening Report</b>", styles['Title']),
        Spacer(1,18)
    ]

    story.append(Paragraph(f"<b>Name:</b> {name}", styles['Normal']))
    story.append(Spacer(1,12))

    risk_val = results.get('risk_index', 'N/A')
    story.append(Paragraph(f"<b>Overall Risk Index (0-100): <font color='red'>{risk_val}</font></b>", styles['Heading2']))
    story.append(Paragraph(descriptions['risk_index'], styleSmall))
    story.append(Spacer(1,20))

    story.append(Paragraph("<b>Individual Test Scores (0-100, Higher is Better):</b>", styles['Heading3']))

    if 'subscores' in results and isinstance(results['subscores'], dict):
        for k, v in results['subscores'].items():
            desc = descriptions.get(k, 'Score out of 100.')
            story.append(Paragraph(f"<b>{k.replace('_', ' ')}:</b> {v}", styles['Normal']))
            story.append(Paragraph(f"<i>({desc})</i>", styleSmall))
            story.append(Spacer(1,6))

    story.append(Spacer(1,18))
    story.append(Paragraph("<b>Raw Metrics:</b>", styles['Heading3']))

    if 'raw_metrics' in results and isinstance(results['raw_metrics'], dict):
        for k, v in results['raw_metrics'].items():
            story.append(Paragraph(f"<b>{k.replace('_', ' ')}:</b> {v}", styles['Normal']))
            story.append(Spacer(1,4))

    story.append(Spacer(1,30))
    story.append(Paragraph("<i>Disclaimer: This tool is for informational screening purposes only and is not a medical diagnostic device.</i>", styleSmall))

    doc.build(story)
    tmp.seek(0)

    safe_name = name.replace(' ', '_').replace('.', '').lower() or 'neurotrack_report'
    download_filename = f'neurotrack_Report_{safe_name}.pdf'

    return send_file(tmp, download_name=download_filename, as_attachment=True, mimetype='application/pdf')

# Vercel serverless compatibility
app = app